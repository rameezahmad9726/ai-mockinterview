"""
Fill FYP thesis presentation template with Interveux project data.
Reads from: d:\FYP (Thesis) - Presentation Template.pptx
Writes to: d:\FYP (Thesis) - Interveux_Filled.pptx
"""
from pathlib import Path
from pptx import Presentation

INPUT_PATH = Path(r"d:\FYP (Thesis) - Presentation Template.pptx")
OUTPUT_PATH = Path(r"d:\FYP (Thesis) - Interveux_Filled.pptx")

SLIDE_CONTENT = {
    1: {
        "title": "Interveux: AI-Powered Mock Interview Platform",
        "subtitle": "Student Name: [Your Name]\nRoll No: [Your Roll No]\nSupervisor: [Supervisor Name]\nCo-supervisor: [Co-supervisor Name]\nDepartment of Information Technology",
    },
    2: """Background of the Study:
• Growing need for interview preparation and objective feedback for job seekers.
• Manual mock interviews are time-consuming and lack consistent, data-driven evaluation.
• AI can combine speech recognition, emotion analysis, and body language to give structured feedback.

This study focuses on building an integrated platform (Interveux) that:
- Parses resumes and generates tailored interview questions
- Conducts live TTS-based interviews with optional video recording
- Analyzes uploaded interview videos for emotion, body language, speech clarity, and confidence
- Produces detailed reports to help candidates improve.""",
    3: """Significance/Contribution:
• Provides a single end-to-end pipeline: resume → questions → live interview → video analysis → report.
• Uses multiple AI models (Whisper, facial emotion classifier, MediaPipe Pose, GPT-4o-mini) in one system.
• Delivers numeric scores (clarity, confidence, posture, movement) and narrative feedback (tone, improvement tips).
• Demonstrates efficient design: parallel video analysis, model reuse, and combined LLM calls to reduce latency and cost.""",
    4: """Research Objectives:
• To design and implement an AI-driven mock interview system that generates resume-based questions and conducts TTS-led interviews.
• To integrate speech-to-text (Whisper), facial emotion recognition, and body pose analysis for multimodal evaluation.
• To derive and present clarity, confidence, posture, and gesture metrics with actionable improvement tips.
• To optimize the pipeline for speed (parallel processing, model caching, batch TTS) while maintaining accuracy.""",
    5: """Research Questions:
• How can resume text be used to automatically generate relevant technical and behavioral interview questions?
• How accurately can off-the-shelf models (Whisper, emotion classifier, pose estimation) evaluate interview performance from video?
• How can clarity and confidence be quantified from speech (e.g., filler words, tone analysis) and aligned in the final report?
• What pipeline design (parallelization, model reuse) makes video analysis feasible for typical hardware?""",
    6: """Research Framework/Model and Hypothesis:
• Framework: Resume upload → question generation (GPT-4o-mini) → Live interview (TTS) → Video upload → Parallel analysis (speech + emotion + body) → Report (scores + tone).

• Models: Whisper (tiny) for STT; dima806/facial_emotions_image_detection for emotion; MediaPipe Pose for body; GPT-4o-mini for questions, diarization, and tone.

• Hypothesis: Combining rule-based metrics (e.g., clarity from filler count) with LLM-based tone analysis and aligning them (e.g., not contradicting high clarity) improves perceived consistency and usefulness of feedback.""",
    7: """Research Methodology:
• Development: FastAPI backend (Python), React frontend; OpenAI API for GPT-4o-mini and TTS; open-source Whisper, Hugging Face emotion model, MediaPipe.
• Video processing: Extract frames (1 FPS) and audio in parallel; run speech, emotion, and body analysis in parallel; load Whisper on main thread to avoid failures.
• Metrics: Clarity = f(filler words); posture = shoulder alignment; movement = hand displacement; confidence/tone = LLM over interviewee transcript.
• Validation: Functional testing of full flow; comparison of model-reported accuracy (e.g., Whisper WER, emotion precision) with literature.""",
}


def add_paragraphs(tf, text: str):
    lines = [t.strip() for t in text.split("\n") if t.strip()]
    if not lines:
        return
    for p in tf.paragraphs:
        p.clear()
    if not tf.paragraphs:
        return
    tf.paragraphs[0].text = lines[0]
    for line in lines[1:]:
        tf.add_paragraph().text = line


def main():
    if not INPUT_PATH.exists():
        print("Input file not found:", INPUT_PATH)
        return
    prs = Presentation(str(INPUT_PATH))
    for slide_idx in range(min(7, len(prs.slides))):
        one_indexed = slide_idx + 1
        content = SLIDE_CONTENT.get(one_indexed)
        if not content:
            continue
        slide = prs.slides[slide_idx]
        if one_indexed == 1:
            for shape in slide.shapes:
                if not hasattr(shape, "text_frame"):
                    continue
                try:
                    ph_type = str(shape.placeholder_format.type)
                except Exception:
                    continue
                if "CENTER" in ph_type and "TITLE" in ph_type:
                    shape.text_frame.paragraphs[0].text = content.get("title", "")
                elif "SUBTITLE" in ph_type:
                    add_paragraphs(shape.text_frame, content.get("subtitle", ""))
            continue
        body_text = content if isinstance(content, str) else content.get("body", "")
        if not body_text:
            continue
        filled = False
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame"):
                continue
            try:
                if shape.is_placeholder and shape.placeholder_format.type == 7:
                    add_paragraphs(shape.text_frame, body_text)
                    filled = True
                    break
            except Exception:
                pass
        if not filled:
            for shape in slide.shapes:
                if not hasattr(shape, "text_frame"):
                    continue
                try:
                    pt = shape.placeholder_format.type
                    if pt in (1, 13):
                        continue
                    add_paragraphs(shape.text_frame, body_text)
                    break
                except Exception:
                    pass
    prs.save(str(OUTPUT_PATH))
    print("Saved:", OUTPUT_PATH)


if __name__ == "__main__":
    main()
